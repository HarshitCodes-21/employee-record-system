import tkinter as tk
from tkinter import ttk, messagebox
from openpyxl import Workbook, load_workbook
import os

# ---------------- FILE SETUP ----------------

file_name = "accounts.xlsx"

if not os.path.exists(file_name):

    wb = Workbook()
    ws = wb.active
    ws.title = "Accounts"

    ws.append(["name", "phone", "salary", "place"])

    wb.save(file_name)

# ---------------- FUNCTIONS ----------------

selected_item = None


def clear_entries():

    entry1.delete(0, tk.END)
    entry2.delete(0, tk.END)
    entry3.delete(0, tk.END)
    entry4.delete(0, tk.END)


# ---------------- SHOW DATA ----------------

def show_data(data=None):

    for row in tree.get_children():
        tree.delete(row)

    if data is not None:

        for row in data:
            tree.insert("", tk.END, values=row)

        return

    wb = load_workbook(file_name)
    ws = wb.active

    for row in ws.iter_rows(min_row=2, values_only=True):
        tree.insert("", tk.END, values=row)


# ---------------- LOAD PLACES ----------------

def load_places():

    wb = load_workbook(file_name)
    ws = wb.active

    places = []

    for row in ws.iter_rows(min_row=2, values_only=True):

        place = row[3]

        if place not in places:
            places.append(place)

    place_combo["values"] = places


# ---------------- LOAD SALARY RANGES ----------------

def load_salary_ranges():

    wb = load_workbook(file_name)
    ws = wb.active

    salary_ranges = ["any"]

    salaries = []

    # collect salaries

    for row in ws.iter_rows(min_row=2, values_only=True):

        salary = int(row[2])
        salaries.append(salary)

    # fixed salary ranges

    ranges = [
        (0, 10000),
        (10000, 50000),
        (50000, 100000),
        (100000, 150000),
        (150000, 200000),
        (200000, 250000),
        (250000, 300000),
        (300000, 350000),
        (350000, 400000),
        (400000, 450000),
        (450000, 500000)
    ]

    # add ranges only if employees exist

    for start, end in ranges:

        for salary in salaries:

            if start <= salary <= end:

                range_name = f"{start}-{end}"

                if range_name not in salary_ranges:
                    salary_ranges.append(range_name)

                break

    salary_combo["values"] = salary_ranges


# ---------------- SAVE DATA ----------------

def save_data():

    name = entry1.get().lower()
    phone = entry2.get().lower()
    salary = entry3.get().lower()
    place = entry4.get().lower()

    if name == "" or phone == "" or salary == "" or place == "":
        messagebox.showwarning("warning", "fill all details")
        return

    wb = load_workbook(file_name)
    ws = wb.active

    ws.append([name, phone, salary, place])

    wb.save(file_name)

    # add place automatically

    current_places = list(place_combo["values"])

    if place not in current_places:
        current_places.append(place)
        place_combo["values"] = current_places

    load_salary_ranges()

    messagebox.showinfo("success", "data saved")

    clear_entries()
    show_data()


# ---------------- SELECT RECORD ----------------

def select_record(event):

    global selected_item

    selected_item = tree.focus()

    values = tree.item(selected_item, "values")

    if values:

        clear_entries()

        entry1.insert(0, values[0])
        entry2.insert(0, values[1])
        entry3.insert(0, values[2])
        entry4.insert(0, values[3])


# ---------------- UPDATE DATA ----------------

def update_data():

    global selected_item

    if not selected_item:
        messagebox.showwarning("warning", "select a record")
        return

    new_values = (
        entry1.get().lower(),
        entry2.get().lower(),
        entry3.get().lower(),
        entry4.get().lower()
    )

    tree.item(selected_item, values=new_values)

    wb = load_workbook(file_name)
    ws = wb.active

    ws.delete_rows(2, ws.max_row)

    for item in tree.get_children():

        row = tree.item(item)["values"]
        ws.append(row)

    wb.save(file_name)

    load_salary_ranges()

    messagebox.showinfo("success", "data updated")

    clear_entries()
    show_data()


# ---------------- DELETE DATA ----------------

def delete_data():

    selected = tree.focus()

    if not selected:
        messagebox.showwarning("warning", "select a record")
        return

    tree.delete(selected)

    wb = load_workbook(file_name)
    ws = wb.active

    ws.delete_rows(2, ws.max_row)

    for item in tree.get_children():

        row = tree.item(item)["values"]
        ws.append(row)

    wb.save(file_name)

    load_salary_ranges()

    messagebox.showinfo("deleted", "record deleted")

    show_data()


# ---------------- FILTER DATA ----------------

def filter_data():

    selected_place = place_combo.get().lower()
    selected_salary = salary_combo.get()

    wb = load_workbook(file_name)
    ws = wb.active

    filtered = []

    for row in ws.iter_rows(min_row=2, values_only=True):

        salary = int(row[2])
        place = row[3]

        # place filter

        if place != selected_place:
            continue

        # any salary

        if selected_salary == "any":

            filtered.append(row)

        else:

            split_range = selected_salary.split("-")

            start = int(split_range[0])
            end = int(split_range[1])

            if start <= salary <= end:
                filtered.append(row)

    if len(filtered) == 0:
        messagebox.showinfo("result", "0 members found")

    show_data(filtered)


# ---------------- SHOW ALL ----------------

def show_all():
    show_data()


# ---------------- GUI ----------------

root = tk.Tk()
root.title("accounts system")
root.geometry("1150x650")

# labels

tk.Label(root, text="enter name").place(x=50, y=30)
tk.Label(root, text="enter phone").place(x=50, y=70)
tk.Label(root, text="enter salary").place(x=50, y=110)
tk.Label(root, text="enter place").place(x=50, y=150)

# entries

entry1 = tk.Entry(root, width=30)
entry2 = tk.Entry(root, width=30)
entry3 = tk.Entry(root, width=30)
entry4 = tk.Entry(root, width=30)

entry1.place(x=170, y=30)
entry2.place(x=170, y=70)
entry3.place(x=170, y=110)
entry4.place(x=170, y=150)

# clear buttons

tk.Button(root, text="clear",
          command=lambda: entry1.delete(0, tk.END)
          ).place(x=380, y=28)

tk.Button(root, text="clear",
          command=lambda: entry2.delete(0, tk.END)
          ).place(x=380, y=68)

tk.Button(root, text="clear",
          command=lambda: entry3.delete(0, tk.END)
          ).place(x=380, y=108)

tk.Button(root, text="clear",
          command=lambda: entry4.delete(0, tk.END)
          ).place(x=380, y=148)

# buttons

tk.Button(root,
          text="save",
          width=10,
          bg="lightblue",
          command=save_data).place(x=30, y=220)

tk.Button(root,
          text="update",
          width=10,
          bg="lightgreen",
          command=update_data).place(x=140, y=220)

tk.Button(root,
          text="delete",
          width=10,
          bg="tomato",
          command=delete_data).place(x=250, y=220)

# ---------------- PLACE FILTER ----------------

tk.Label(root, text="select place").place(x=50, y=300)

place_combo = ttk.Combobox(root)

place_combo.place(x=170, y=300)

# ---------------- SALARY FILTER ----------------

tk.Label(root, text="salary range").place(x=50, y=350)

salary_combo = ttk.Combobox(root)

salary_combo.place(x=170, y=350)

# ---------------- FILTER BUTTON ----------------

tk.Button(root,
          text="filter data",
          bg="orange",
          command=filter_data).place(x=180, y=400)

# ---------------- SHOW ALL BUTTON ----------------

tk.Button(root,
          text="show all",
          bg="yellow",
          command=show_all).place(x=300, y=400)

# ---------------- TABLE ----------------

columns = ("name", "phone", "salary", "place")

tree = ttk.Treeview(root,
                    columns=columns,
                    show="headings")

for col in columns:

    tree.heading(col, text=col)
    tree.column(col, width=180)

tree.place(x=500, y=30, width=600, height=500)

# select row

tree.bind("<ButtonRelease-1>", select_record)

# load data

load_places()
load_salary_ranges()
show_data()

root.mainloop()
